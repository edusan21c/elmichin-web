import os
import re
import openpyxl

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RUTA_ENTRADA = os.path.join(BASE_DIR, 'data', 'Inventario_Maestro_Final.xlsx')
RUTA_SALIDA = os.path.join(BASE_DIR, 'data', 'Inventario_Maestro_Sin_Prefijos.xlsx')

def limpiar_prefijos():
    print("--- LIMPIANDO PREFIJOS DE LOS NOMBRES ---")

    if not os.path.exists(RUTA_ENTRADA):
        print(f"ERROR: No se encontró {RUTA_ENTRADA}")
        return

    # Cargar el archivo conservando todo el formato
    wb = openpyxl.load_workbook(RUTA_ENTRADA)
    ws = wb.active

    headers = [cell.value for cell in ws[1]]
    col_nombre_idx = headers.index('Nombre Producto') + 1

    # Patrón para los prefijos. Busca al inicio de la celda (^), 
    # que empiece con Un. U. D. Ag. B. (ignorando mayúsculas/minúsculas) y borra los espacios extra.
    patron = re.compile(r'^(Un\.|U\.|D\.|Ag\.|B\.)\s*', re.IGNORECASE)

    modificados = 0

    for row in range(2, ws.max_row + 1):
        cell = ws.cell(row=row, column=col_nombre_idx)
        nombre_original = str(cell.value).strip()
        
        # Aplicamos el reemplazo
        nombre_limpio = patron.sub('', nombre_original).strip()
        
        # Si el nombre cambió, lo actualizamos
        if nombre_limpio != nombre_original:
            cell.value = nombre_limpio
            modificados += 1

    wb.save(RUTA_SALIDA)
    print(f"\n--- TERMINADO ---")
    print(f"Nombres limpiados: {modificados}")
    print(f"Archivo guardado como: {RUTA_SALIDA}")

if __name__ == '__main__':
    limpiar_prefijos()