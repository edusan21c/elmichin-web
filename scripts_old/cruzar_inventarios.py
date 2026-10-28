import os
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import unicodedata

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RUTA_MICHIN = os.path.join(BASE_DIR, 'data', 'Inventario_Maestro_Final.xlsx')
RUTA_LOGGRO = os.path.join(BASE_DIR, 'data', 'Loggro_Ordenado.xlsx')
RUTA_SALIDA = os.path.join(BASE_DIR, 'data', 'Reporte_Cruce_Inventarios.xlsx')

def limpiar_texto(txt):
    if pd.isna(txt): return ""
    txt = str(txt).strip().upper()
    txt = ''.join(c for c in unicodedata.normalize('NFD', txt) if unicodedata.category(c) != 'Mn')
    return txt

def obtener_codigo_limpio(val):
    if pd.isna(val): return ""
    val_str = str(val).strip()
    if val_str.endswith('.0'): val_str = val_str[:-2]
    return val_str

def estilizar_hoja(ws, titulo_color):
    # Estilo para los encabezados
    header_fill = PatternFill(start_color=titulo_color, end_color=titulo_color, fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Auto-ajustar ancho de columnas
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2)
        ws.column_dimensions[column].width = adjusted_width
        
    # Congelar la primera fila y activar filtros
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

def cruzar():
    print("--- GENERANDO REPORTE DE CRUCE DE INVENTARIOS ---")

    if not os.path.exists(RUTA_MICHIN) or not os.path.exists(RUTA_LOGGRO):
        print("ERROR: Faltan archivos base.")
        return

    df_michin = pd.read_excel(RUTA_MICHIN)
    df_loggro = pd.read_excel(RUTA_LOGGRO)

    # 1. Limpiar Michín
    michin_codigos = set()
    michin_nombres = set()
    michin_limpio = []
    
    for _, r in df_michin.iterrows():
        cod = obtener_codigo_limpio(r.get('Código Barras'))
        nom = limpiar_texto(r.get('Nombre Producto'))
        if cod: michin_codigos.add(cod)
        if nom: michin_nombres.add(nom)
        michin_limpio.append({'Nombre': r.get('Nombre Producto'), 'Codigo': cod, 'Categoria': r.get('Categoría')})

    # 2. Limpiar Loggro y quitar repetidos
    loggro_codigos = set()
    loggro_nombres = set()
    loggro_limpio = []
    loggro_vistos = set() # Para evitar repetidos en el reporte

    for _, r in df_loggro.iterrows():
        cod = obtener_codigo_limpio(r.get('Ítem'))
        nom = limpiar_texto(r.get('Descripción'))
        
        if not nom: continue
        if cod: loggro_codigos.add(cod)
        loggro_nombres.add(nom)
        
        # Evitar duplicados exactos en Loggro
        firma = f"{cod}|{nom}"
        if firma not in loggro_vistos:
            loggro_vistos.add(firma)
            loggro_limpio.append({'Nombre': r.get('Descripción'), 'Codigo': cod, 'Categoria': r.get('Categoría')})

    # 3. Cruzar datos (Los que están en Loggro pero NO en Michín)
    faltan_en_michin = []
    for item in loggro_limpio:
        cod = item['Codigo']
        nom = limpiar_texto(item['Nombre'])
        
        # Si el código no está en Michín Y el nombre tampoco, entonces falta.
        # (Usamos texto exacto para no cruzar cosas equivocadas)
        if (cod not in michin_codigos) and (nom not in michin_nombres):
            faltan_en_michin.append(item)

    # 4. Cruzar datos (Los que están en Michín pero NO en Loggro)
    faltan_en_loggro = []
    for item in michin_limpio:
        cod = item['Codigo']
        nom = limpiar_texto(item['Nombre'])
        
        if (cod not in loggro_codigos) and (nom not in loggro_nombres):
            faltan_en_loggro.append(item)

    # Convertir a DataFrames y ordenar alfabéticamente
    df_faltan_michin = pd.DataFrame(faltan_en_michin).sort_values(by='Nombre')
    df_faltan_loggro = pd.DataFrame(faltan_en_loggro).sort_values(by='Nombre')

    # Renombrar columnas para que se vean bonitas
    df_faltan_michin.columns = ['Descripción en Loggro', 'Código en Loggro', 'Categoría en Loggro']
    df_faltan_loggro.columns = ['Producto en Michín', 'Código en Michín', 'Categoría en Michín']

    # 5. Guardar en Excel con diseño
    with pd.ExcelWriter(RUTA_SALIDA, engine='openpyxl') as writer:
        df_faltan_michin.to_excel(writer, sheet_name='Faltan en Michin', index=False)
        df_faltan_loggro.to_excel(writer, sheet_name='Faltan en Loggro', index=False)

    # 6. Aplicar estilos (colores a los encabezados)
    wb = openpyxl.load_workbook(RUTA_SALIDA)
    
    # Hoja 1: Faltan en Michín (Color Azul)
    ws1 = wb['Faltan en Michin']
    estilizar_hoja(ws1, "1F4E78") 
    
    # Hoja 2: Faltan en Loggro (Color Verde)
    ws2 = wb['Faltan en Loggro']
    estilizar_hoja(ws2, "375623")

    wb.save(RUTA_SALIDA)

    print(f"\n--- TERMINADO ---")
    print(f"Productos en Loggro que NO tienes en Michín: {len(df_faltan_michin)}")
    print(f"Productos en Michín que NO están en Loggro: {len(df_faltan_loggro)}")
    print(f"Archivo guardado en: {RUTA_SALIDA}")

if __name__ == '__main__':
    cruzar()