import os
import shutil
import openpyxl

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RUTA_MAESTRO = os.path.join(BASE_DIR, 'data', 'Inventario_Maestro_Completo.xlsx')
RUTA_SALIDA = os.path.join(BASE_DIR, 'data', 'Inventario_Maestro_Final.xlsx')

def aplicar():
    print("--- LIMPIEZA SEGURA: VACIAR CELDAS SIN CODIGO ---")

    if not os.path.exists(RUTA_MAESTRO):
        print("ERROR: Falta el archivo original.")
        return

    # 1. Copiamos el original (el que NO tiene errores de Loggro)
    shutil.copyfile(RUTA_MAESTRO, RUTA_SALIDA)

    # 2. Abrimos el archivo copiado
    wb = openpyxl.load_workbook(RUTA_SALIDA)
    ws = wb.active

    headers = [cell.value for cell in ws[1]]
    col_cod = headers.index('Código Barras') + 1

    vaciados = 0
    conservados = 0

    # 3. Recorremos todas las filas
    for row in range(2, ws.max_row + 1):
        cell = ws.cell(row=row, column=col_cod)
        
        # Forzamos a que toda la columna sea formato Texto
        cell.number_format = '@'

        val = cell.value
        
        # Limpiamos el valor por si viene como número flotante de Excel
        if val is not None:
            if isinstance(val, float) and val.is_integer():
                val_str = str(int(val))
            else:
                val_str = str(val).strip()
        else:
            val_str = ""

        # LA REGLA DE ORO: Si no son puros números, se borra la celda
        if not val_str.isdigit():
            cell.value = None  # Deja la celda en blanco
            vaciados += 1
        else:
            cell.value = val_str # Conserva el código real y lo deja como texto
            conservados += 1

    # 4. Guardamos el archivo final
    wb.save(RUTA_SALIDA)
    
    print(f"\n--- TERMINADO ---")
    print(f"Codigos reales conservados: {conservados}")
    print(f"Celdas vaciadas (basura o vacias): {vaciados}")
    print(f"Archivo limpio guardado en: {RUTA_SALIDA}")

if __name__ == '__main__':
    aplicar()