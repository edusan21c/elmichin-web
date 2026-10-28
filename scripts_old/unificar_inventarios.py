import os
import pandas as pd
import openpyxl
import re
import unicodedata
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RUTA_MICHIN = os.path.join(BASE_DIR, 'data', 'Inventario_Maestro_Final.xlsx')
RUTA_LOGGRO = os.path.join(BASE_DIR, 'data', 'Loggro_Ordenado.xlsx')
RUTA_SALIDA = os.path.join(BASE_DIR, 'data', 'Inventario_Unificado.xlsx')

def limpiar_texto(txt):
    if pd.isna(txt): return ""
    txt = str(txt).strip().upper()
    txt = ''.join(c for c in unicodedata.normalize('NFD', txt) if unicodedata.category(c) != 'Mn')
    return txt

def obtener_codigo_limpio(val):
    if pd.isna(val): return ""
    val_str = str(val).strip()
    if val_str.endswith('.0'): val_str = val_str[:-2]
    return val_str

def formatear_nombre(txt):
    if pd.isna(txt): return ""
    # Convierte a Title Case (Primera Mayúscula, resto minúsculas)
    return str(txt).strip().title()

def obtener_palabras_clave(txt):
    # Extrae palabras de 4 letras o más para buscar coincidencias
    if pd.isna(txt): return set()
    txt = limpiar_texto(txt)
    return set(re.findall(r'\b[A-Z]{4,}\b', txt))

def unificar():
    print("--- UNIFICANDO INVENTARIOS (CON DISEÑO Y COLORES) ---")

    if not os.path.exists(RUTA_MICHIN) or not os.path.exists(RUTA_LOGGRO):
        print("ERROR: Faltan archivos base.")
        return

    df_michin = pd.read_excel(RUTA_MICHIN)
    
    # Formatear nombres de Michín
    df_michin['Nombre Producto'] = df_michin['Nombre Producto'].apply(formatear_nombre)
    df_michin['_EsNuevo'] = False
    df_michin['_PosibleDuplicado'] = False

    michin_codigos = set()
    michin_nombres = set()
    michin_palabras = set()
    
    for _, r in df_michin.iterrows():
        cod = obtener_codigo_limpio(r.get('Código Barras'))
        nom = limpiar_texto(r.get('Nombre Producto'))
        if cod: michin_codigos.add(cod)
        if nom: michin_nombres.add(nom)
        michin_palabras.update(obtener_palabras_clave(nom))

    df_loggro = pd.read_excel(RUTA_LOGGRO)

    nuevos_productos = []
    loggro_vistos = set()

    for _, r in df_loggro.iterrows():
        cod = obtener_codigo_limpio(r.get('Ítem'))
        nom_orig = r.get('Descripción')
        nom_limpio = limpiar_texto(nom_orig)
        
        if not nom_limpio: continue

        firma = f"{cod}|{nom_limpio}"
        if firma in loggro_vistos: continue
        loggro_vistos.add(firma)

        # Si ya existe en Michín, lo ignoramos
        if (cod and cod in michin_codigos) or (nom_limpio in michin_nombres):
            continue 

        # Comprobar si comparte palabras clave con Michín (Para pintarlo de otro color)
        palabras_loggro = obtener_palabras_clave(nom_limpio)
        es_posible_duplicado = bool(palabras_loggro.intersection(michin_palabras))

        nuevo = {
            'ID (NO TOCAR)': None,
            'Global (NO TOCAR)': None,
            'Nombre Producto': formatear_nombre(nom_orig), # Formato Título
            'Categoría': None,
            'Código Barras': cod,
            'T1 Stock': 0, 'T1 Costo': 0, 'T1 % Gan': 0, 'T1 Precio Base': r.get('Precio Vta.', 0),
            'T1 Cond1': 0, 'T1 Precio 1': 0, 'T1 Cond2': 0, 'T1 Precio 2': 0, 'T1 Cond3': 0, 'T1 Precio 3': 0,
            'T2 Stock': 0, 'T2 Costo': 0, 'T2 % Gan': 0, 'T2 Precio Base': r.get('Precio Vta.', 0),
            'T2 Cond1': 0, 'T2 Precio 1': 0, 'T2 Cond2': 0, 'T2 Precio 2': 0, 'T2 Cond3': 0, 'T2 Precio 3': 0,
            'Tiene Foto?': 'NO',
            '_EsNuevo': True,
            '_PosibleDuplicado': es_posible_duplicado
        }
        nuevos_productos.append(nuevo)

    print(f"Productos nuevos importados de Loggro: {len(nuevos_productos)}")

    df_nuevos = pd.DataFrame(nuevos_productos)
    df_final = pd.concat([df_michin, df_nuevos], ignore_index=True)

    # Ordenar Alfabéticamente
    df_final = df_final.sort_values(by='Nombre Producto', ascending=True)

    df_final.to_excel(RUTA_SALIDA, index=False)

    # ---------------------------------------------------------
    # DISEÑO EN EXCEL
    # ---------------------------------------------------------
    wb = openpyxl.load_workbook(RUTA_SALIDA)
    ws = wb.active

    # Estilo encabezado
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Colores para las filas
    fill_amarillo_claro = PatternFill(start_color="FFFF99", end_color="FFFF99", fill_type="solid") # Nuevos
    fill_amarillo_fuerte = PatternFill(start_color="FFD966", end_color="FFD966", fill_type="solid") # Posibles Duplicados

    headers = [cell.value for cell in ws[1]]
    col_cod_idx = headers.index('Código Barras') + 1
    col_nuevo_idx = headers.index('_EsNuevo') + 1
    col_dup_idx = headers.index('_PosibleDuplicado') + 1

    for row in range(2, ws.max_row + 1):
        # Formato de texto para códigos
        cell_cod = ws.cell(row=row, column=col_cod_idx)
        cell_cod.number_format = '@'
        if cell_cod.value:
            cell_cod.value = str(cell_cod.value).strip()

        es_nuevo = ws.cell(row=row, column=col_nuevo_idx).value
        es_dup = ws.cell(row=row, column=col_dup_idx).value

        if es_nuevo:
            color = fill_amarillo_fuerte if es_dup else fill_amarillo_claro
            for col_idx in range(1, col_nuevo_idx): # Pintar hasta antes de columnas ocultas
                ws.cell(row=row, column=col_idx).fill = color

    # Ocultar columnas de control y auto-ajustar
    ws.column_dimensions[get_column_letter(col_nuevo_idx)].hidden = True
    ws.column_dimensions[get_column_letter(col_dup_idx)].hidden = True

    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except: pass
        ws.column_dimensions[column].width = min(max_length + 2, 50)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(col_nuevo_idx-1)}{ws.max_row}"

    wb.save(RUTA_SALIDA)

    print(f"\n--- TERMINADO ---")
    print(f"Archivo guardado en: {RUTA_SALIDA}")
    print(f"Total de productos: {len(df_final)}")

if __name__ == '__main__':
    unificar()