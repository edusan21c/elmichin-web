import os
import re
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import unicodedata

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RUTA_MICHIN = os.path.join(BASE_DIR, 'data', 'Inventario_Maestro_Coloreado.xlsx')
RUTA_LOGGRO = os.path.join(BASE_DIR, 'data', 'Loggro_Ordenado.xlsx')
RUTA_SALIDA = os.path.join(BASE_DIR, 'data', 'Inventario_Maestro_Unificado_Final.xlsx')

def limpiar_texto(txt):
    if pd.isna(txt): return ""
    txt = str(txt).strip().upper()
    txt = ''.join(c for c in unicodedata.normalize('NFD', txt) if unicodedata.category(c) != 'Mn')
    return txt

def obtener_codigo_limpio(val):
    if pd.isna(val): return ""
    val_str = str(val).strip()
    if val_str.endswith('.0'): val_str = val_str[:-2]
    return val_str

def cruzar():
    print("--- CRUCE DEFINITIVO CON LOGGRO (NUEVOS EN FUCSIA) ---")

    if not os.path.exists(RUTA_MICHIN) or not os.path.exists(RUTA_LOGGRO):
        print("ERROR: Faltan archivos base.")
        return

    # 1. Leer Michín (Conservando colores originales)
    wb_michin = openpyxl.load_workbook(RUTA_MICHIN)
    ws_michin = wb_michin.active
    headers = [cell.value for cell in ws_michin[1]]
    
    col_nombre_idx = headers.index('Nombre Producto') + 1

    michin_datos = []
    michin_nombres_limpios = set()
    michin_codigos = set()

    for row in range(2, ws_michin.max_row + 1):
        fila_data = {}
        for col_idx, header in enumerate(headers, start=1):
            fila_data[header] = ws_michin.cell(row=row, column=col_idx).value
        
        nom_original = fila_data.get('Nombre Producto')
        if not nom_original: continue

        # Extraer el color de la celda del nombre
        color_celda = ws_michin.cell(row=row, column=col_nombre_idx).fill.start_color.rgb
        if color_celda == '00000000':
            color_celda = None
            
        fila_data['_color'] = color_celda
        
        michin_datos.append(fila_data)
        michin_nombres_limpios.add(limpiar_texto(nom_original))
        cod = obtener_codigo_limpio(fila_data.get('Código Barras'))
        if cod: michin_codigos.add(cod)

    # 2. Leer Loggro y buscar los que faltan
    df_loggro = pd.read_excel(RUTA_LOGGRO)
    
    nuevos_loggro = []
    loggro_vistos = set()

    for _, r in df_loggro.iterrows():
        cod = obtener_codigo_limpio(r.get('Ítem'))
        nom_orig = r.get('Descripción')
        nom_limpio = limpiar_texto(nom_orig)

        if not nom_limpio: continue
        
        firma = f"{cod}|{nom_limpio}"
        if firma in loggro_vistos: continue
        loggro_vistos.add(firma)

        # Si ya existe en Michín, no lo agregamos
        if (cod and cod in michin_codigos) or (nom_limpio in michin_nombres_limpios):
            continue

        nuevo_producto = {header: None for header in headers}
        nuevo_producto['Nombre Producto'] = str(nom_orig).strip().title()
        nuevo_producto['Categoría'] = None
        nuevo_producto['Código Barras'] = cod if cod else None
        nuevo_producto['T1 Stock'] = 0
        nuevo_producto['T1 Costo'] = 0
        nuevo_producto['T1 % Gan'] = 0
        nuevo_producto['T1 Precio Base'] = r.get('Precio Vta.', 0)
        nuevo_producto['T2 Stock'] = 0
        nuevo_producto['T2 Costo'] = 0
        nuevo_producto['T2 % Gan'] = 0
        nuevo_producto['T2 Precio Base'] = r.get('Precio Vta.', 0)
        nuevo_producto['Tiene Foto?'] = 'NO'
        
        # 🔥 AQUI ESTA EL COLOR FUCSIA BRILLANTE 🔥
        nuevo_producto['_color'] = 'FFFF00FF' 
        
        nuevos_loggro.append(nuevo_producto)

    print(f"Productos nuevos de Loggro agregados: {len(nuevos_loggro)}")

    # 3. Unir y Ordenar Alfabéticamente
    todos_los_productos = michin_datos + nuevos_loggro
    todos_los_productos.sort(key=lambda x: str(x.get('Nombre Producto')).upper())

    # 4. Guardar en nuevo Excel
    wb_final = openpyxl.Workbook()
    ws_final = wb_final.active
    ws_final.title = "Inventario"

    # Escribir encabezados
    for col_idx, header in enumerate(headers, start=1):
        cell = ws_final.cell(row=1, column=col_idx, value=header)
        cell.fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Escribir datos
    for row_idx, prod in enumerate(todos_los_productos, start=2):
        for col_idx, header in enumerate(headers, start=1):
            cell = ws_final.cell(row=row_idx, column=col_idx, value=prod.get(header))
            
            # Aplicar color (Fucsia para nuevos, original para el resto)
            if prod['_color']:
                cell.fill = PatternFill(start_color=prod['_color'], end_color=prod['_color'], fill_type="solid")
            
            # Formato de texto para el código
            if header == 'Código Barras':
                cell.number_format = '@'
                if cell.value:
                    cell.value = str(cell.value).strip()

    # Auto-ajustar columnas
    for col in ws_final.columns:
        max_length = max((len(str(cell.value)) for cell in col if cell.value), default=0)
        ws_final.column_dimensions[col[0].column_letter].width = min(max_length + 2, 50)

    ws_final.freeze_panes = "A2"
    ws_final.auto_filter.ref = ws_final.dimensions

    wb_final.save(RUTA_SALIDA)
    print(f"\n--- TERMINADO ---")
    print(f"Total de productos en el catálogo final: {len(todos_los_productos)}")
    print(f"Archivo guardado como: {RUTA_SALIDA}")

if __name__ == '__main__':
    cruzar()