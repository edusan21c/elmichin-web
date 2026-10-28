import os
import sys
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.append(BASE_DIR)

from app import create_app
from app.models.producto import Producto, ProductoTienda

app = create_app()
RUTA_EXCEL = os.path.join(BASE_DIR, 'data', 'Inventario_Maestro_Completo.xlsx')

def exportar():
    print("--- GENERANDO EXCEL MAESTRO COMPLETO (T1 y T2) ---")
    with app.app_context():
        productos = Producto.query.order_by(Producto.nombre).all()
        
        datos = []
        for p in productos:
            pres_t1 = ProductoTienda.query.filter_by(producto_id=p.id, tienda_id=1).first()
            pres_t2 = ProductoTienda.query.filter_by(producto_id=p.id, tienda_id=2).first()

            fila = {
                'ID (NO TOCAR)': p.id,
                'Global (NO TOCAR)': p.codigo_global,
                'Nombre Producto': p.nombre,
                'Categoría': p.categoria or '',
                'Código Barras': p.codigo_barras or '',
                
                # --- TIENDA 1 ---
                'T1 Stock': pres_t1.cantidad if pres_t1 else 0,
                'T1 Costo': float(pres_t1.precio_proveedor) if pres_t1 and pres_t1.precio_proveedor else 0,
                'T1 % Gan': float(pres_t1.porcentaje) if pres_t1 and pres_t1.porcentaje else 0,
                'T1 Precio Base': float(pres_t1.precio_venta) if pres_t1 and pres_t1.precio_venta else 0,
                'T1 Cond1': pres_t1.condicion1 if pres_t1 else '',
                'T1 Precio 1': float(pres_t1.precio_venta1) if pres_t1 and pres_t1.precio_venta1 else 0,
                'T1 Cond2': pres_t1.condicion2 if pres_t1 else '',
                'T1 Precio 2': float(pres_t1.precio_venta2) if pres_t1 and pres_t1.precio_venta2 else 0,
                'T1 Cond3': pres_t1.condicion3 if pres_t1 else '',
                'T1 Precio 3': float(pres_t1.precio_venta3) if pres_t1 and pres_t1.precio_venta3 else 0,

                # --- TIENDA 2 ---
                'T2 Stock': pres_t2.cantidad if pres_t2 else 0,
                'T2 Costo': float(pres_t2.precio_proveedor) if pres_t2 and pres_t2.precio_proveedor else 0,
                'T2 % Gan': float(pres_t2.porcentaje) if pres_t2 and pres_t2.porcentaje else 0,
                'T2 Precio Base': float(pres_t2.precio_venta) if pres_t2 and pres_t2.precio_venta else 0,
                'T2 Cond1': pres_t2.condicion1 if pres_t2 else '',
                'T2 Precio 1': float(pres_t2.precio_venta1) if pres_t2 and pres_t2.precio_venta1 else 0,
                'T2 Cond2': pres_t2.condicion2 if pres_t2 else '',
                'T2 Precio 2': float(pres_t2.precio_venta2) if pres_t2 and pres_t2.precio_venta2 else 0,
                'T2 Cond3': pres_t2.condicion3 if pres_t2 else '',
                'T2 Precio 3': float(pres_t2.precio_venta3) if pres_t2 and pres_t2.precio_venta3 else 0,

                'Tiene Foto?': 'SÍ' if p.imagen else 'NO'
            }
            datos.append(fila)

        df = pd.DataFrame(datos)
        
        with pd.ExcelWriter(RUTA_EXCEL, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name='Inventario', index=False)
            workbook = writer.book
            worksheet = writer.sheets['Inventario']

            # Estilos
            header_font = Font(bold=True, color='FFFFFF')
            header_fill_t1 = PatternFill(start_color='0D6EFD', end_color='0D6EFD', fill_type='solid') # Azul T1
            header_fill_t2 = PatternFill(start_color='198754', end_color='198754', fill_type='solid') # Verde T2
            header_fill_general = PatternFill(start_color='6C757D', end_color='6C757D', fill_type='solid') # Gris general

            for col_num, value in enumerate(df.columns.values):
                cell = worksheet.cell(row=1, column=col_num + 1)
                cell.font = header_font
                if 'T1' in value:
                    cell.fill = header_fill_t1
                elif 'T2' in value:
                    cell.fill = header_fill_t2
                else:
                    cell.fill = header_fill_general
                
                cell.alignment = Alignment(horizontal='center', vertical='center')

            for i, col in enumerate(df.columns):
                column_len = max(df[col].astype(str).map(len).max(), len(col)) + 2
                worksheet.column_dimensions[get_column_letter(i + 1)].width = column_len

            worksheet.freeze_panes = 'F2' # Congela las primeras 5 columnas y la fila 1

    print(f"--- EXCEL GENERADO EN: {RUTA_EXCEL} ---")
    os.startfile(RUTA_EXCEL)

if __name__ == '__main__':
    exportar()